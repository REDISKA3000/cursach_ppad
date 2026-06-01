from __future__ import annotations

import json
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Iterable

from sqlalchemy.orm import Session

from app.config import PROJECT_ROOT
from app.schemas.role_relevance import (
    CandidateVacancy,
    FinalRankedVacancy,
    RoleRelevancePipelineResult,
    RoleRelevancePipelineStats,
)
from app.services.role_profile_builder import build_role_profile
from app.services.vacancy_candidate_pool import VacancyCandidatePoolBuilder
from app.services.vacancy_hybrid_retriever import load_job_board_vacancies
from app.services.vacancy_llm_reranker import VacancyLLMReranker

logger = logging.getLogger(__name__)


def run_role_relevance_pipeline(
    db: Session,
    target_role: str,
    limit: int = 300,
    use_llm_reranker: bool = True,
) -> RoleRelevancePipelineResult:
    role_profile = build_role_profile(target_role, use_llm=False)
    vacancies = load_job_board_vacancies(db)
    logger.info("[HybridSearch] total vacancies=%s", len(vacancies))

    pool_builder = VacancyCandidatePoolBuilder()
    candidate_pool = pool_builder.build_candidate_pool(role_profile, vacancies, limit=limit)

    reranker = VacancyLLMReranker()
    rerank_results = reranker.rerank_candidate_pool(
        role_profile,
        candidate_pool,
        batch_size=5,
        use_llm=use_llm_reranker,
    )
    rerank_by_id = {result.vacancy_id: result for result in rerank_results}

    final_results = _build_final_results(candidate_pool, rerank_by_id)
    stats = RoleRelevancePipelineStats(
        total_vacancies=len(vacancies),
        candidate_pool_size=len(candidate_pool),
        core_count=sum(1 for item in final_results if item.market_segment == "core"),
        adjacent_count=sum(1 for item in final_results if item.market_segment == "adjacent"),
        excluded_count=sum(1 for item in final_results if item.market_segment == "excluded"),
    )
    logger.info(
        "[Reranker] core=%s, adjacent=%s, excluded=%s",
        stats.core_count,
        stats.adjacent_count,
        stats.excluded_count,
    )
    return RoleRelevancePipelineResult(
        target_role=target_role,
        role_profile=role_profile,
        stats=stats,
        results=final_results,
    )


def _build_final_results(candidates: list[CandidateVacancy], rerank_by_id: dict) -> list[FinalRankedVacancy]:
    results: list[FinalRankedVacancy] = []
    for candidate in candidates:
        rerank = rerank_by_id.get(candidate.vacancy_id)
        if not rerank:
            continue
        results.append(
            FinalRankedVacancy(
                vacancy_id=candidate.vacancy_id,
                title=candidate.title,
                company=candidate.company,
                url=candidate.url,
                hybrid_retrieval_score=candidate.hybrid_retrieval_score,
                semantic_score=candidate.semantic_score,
                keyword_score=candidate.keyword_score,
                structured_score=candidate.structured_score,
                quality_score=candidate.quality_score,
                rerank_score=rerank.rerank_score,
                confidence=rerank.confidence,
                market_segment=rerank.market_segment,
                reason=rerank.reason,
                matched_signals=candidate.matched_signals,
            )
        )
    segment_order = {"core": 0, "adjacent": 1, "excluded": 2}
    results.sort(
        key=lambda item: (
            segment_order[item.market_segment],
            -item.rerank_score,
            -item.hybrid_retrieval_score,
        )
    )
    return results


def export_role_relevance_result(
    result: RoleRelevancePipelineResult,
    output_dir: str | Path | None = None,
    basename: str | None = None,
) -> tuple[Path, Path]:
    output_path = Path(output_dir or Path(PROJECT_ROOT) / "outputs")
    output_path.mkdir(parents=True, exist_ok=True)
    stem = basename or f"role_relevance_{_slugify(result.target_role)}"
    json_path = output_path / f"{stem}.json"
    xlsx_path = output_path / f"{stem}.xlsx"

    json_path.write_text(result.model_dump_json(indent=2), encoding="utf-8")
    _write_xlsx(result, xlsx_path)
    logger.info("[Export] saved results to %s and %s", json_path, xlsx_path)
    return json_path, xlsx_path


def _slugify(value: str) -> str:
    normalized = re.sub(r"[^a-zA-Zа-яА-ЯёЁ0-9]+", "_", value.lower()).strip("_")
    if "продукт" in normalized and "аналит" in normalized:
        return "product_analyst"
    return normalized or datetime.utcnow().strftime("%Y%m%d_%H%M%S")


def _write_xlsx(result: RoleRelevancePipelineResult, path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to export XLSX. Install requirements.txt first.") from exc

    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    rows = [_vacancy_row(item) for item in result.results]
    _add_sheet(workbook, "all_ranked", rows)
    for segment in ["core", "adjacent", "excluded"]:
        _add_sheet(workbook, segment, [row for row in rows if row["market_segment"] == segment])

    summary = workbook.create_sheet("summary", 0)
    summary_rows = [
        ("target_role", result.target_role),
        ("canonical_role", result.role_profile.canonical_role),
        ("role_family", result.role_profile.role_family),
        ("total_vacancies", result.stats.total_vacancies),
        ("candidate_pool_size", result.stats.candidate_pool_size),
        ("core_count", result.stats.core_count),
        ("adjacent_count", result.stats.adjacent_count),
        ("excluded_count", result.stats.excluded_count),
    ]
    for row_index, row in enumerate(summary_rows, start=1):
        summary.cell(row=row_index, column=1, value=row[0])
        summary.cell(row=row_index, column=2, value=row[1])
    summary["A1"].font = Font(bold=True)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 40

    workbook.save(path)


def _vacancy_row(item: FinalRankedVacancy) -> dict:
    keywords = item.matched_signals.get("keywords") or {}
    structured = item.matched_signals.get("structured") or {}
    return {
        "vacancy_id": item.vacancy_id,
        "market_segment": item.market_segment,
        "rerank_score": item.rerank_score,
        "confidence": item.confidence,
        "hybrid_retrieval_score": item.hybrid_retrieval_score,
        "semantic_score": item.semantic_score,
        "keyword_score": item.keyword_score,
        "structured_score": item.structured_score,
        "quality_score": item.quality_score,
        "title": item.title,
        "company": item.company or "",
        "url": item.url or "",
        "reason": item.reason,
        "matched_synonyms": ", ".join(keywords.get("synonyms") or []),
        "matched_must_have": ", ".join(keywords.get("must_have") or []),
        "matched_related": ", ".join(keywords.get("related") or []),
        "matched_negative": ", ".join(keywords.get("negative") or []),
        "matched_adjacent": ", ".join(keywords.get("adjacent") or []),
        "matched_excluded": ", ".join(keywords.get("excluded") or []),
        "matched_role_family": structured.get("matched_role_family", False),
    }


def _add_sheet(workbook, title: str, rows: Iterable[dict]) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = list(rows)
    headers = [
        "vacancy_id",
        "market_segment",
        "rerank_score",
        "confidence",
        "hybrid_retrieval_score",
        "semantic_score",
        "keyword_score",
        "structured_score",
        "quality_score",
        "title",
        "company",
        "url",
        "reason",
        "matched_synonyms",
        "matched_must_have",
        "matched_related",
        "matched_negative",
        "matched_adjacent",
        "matched_excluded",
        "matched_role_family",
    ]
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "A": 12,
        "B": 15,
        "C": 13,
        "D": 11,
        "E": 20,
        "J": 42,
        "K": 22,
        "L": 48,
        "M": 80,
    }
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = widths.get(letter, min(max(len(header) + 2, 14), 30))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")
